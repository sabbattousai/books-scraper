"""Independent verification PASS 2 (different route to the data + checks the Excel file).

Reads kitaplar.xlsx back (cell values AND hyperlink targets), then crawls the 50
*category listing pages* (not the book pages) and matches every book by its id:
title, price, rating, category and link must agree. Also checks the Excel table
and summary-sheet formulas exist. Writes recheck_categories_report.csv.
"""
import re
import sys
import time
from urllib.parse import urljoin

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

RATING = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}
s = requests.Session()


def get(url):
    for i in range(3):
        try:
            r = s.get(url, timeout=15)
            r.raise_for_status()
            return BeautifulSoup(r.content.decode("utf-8"), "lxml")
        except requests.RequestException:
            time.sleep(2 * (i + 1))
    return None


# ---- what the Excel file really contains ----
wb = load_workbook("kitaplar.xlsx")  # formulas stay as text, hyperlinks readable
ws = wb["Kitaplar"]
problems = []
assert [c.value for c in ws[1]] == ["Kitap ID", "Kitap Adı", "Fiyat", "Puan", "Stok Durumu", "Stok Adedi", "Kategori", "Link", "Kısa Link"]
xl = {}
for r in range(2, ws.max_row + 1):
    row = [c.value for c in ws[r]]
    link_cell = ws.cell(r, 8)
    target = link_cell.hyperlink.target if link_cell.hyperlink else None
    if target != row[7]:
        problems.append((row[7], f"Excel hyperlink target {target} != cell text"))
    xl[row[0]] = dict(title=row[1], price=row[2], rating=row[3], cat=row[6], link=row[7])
print(f"Excel rows read: {len(xl)} | tables: {list(ws.tables)} | sheets: {wb.sheetnames}")
if "Kitaplar" not in ws.tables:
    problems.append(("-", "Excel table missing"))
if not str(wb["Kateqoriya özət"]["C2"].value).startswith("=AVERAGEIFS"):
    problems.append(("-", "summary formulas missing"))

# ---- crawl category listing pages ----
home = get("https://books.toscrape.com/index.html")
cats = [(a.get_text(strip=True), urljoin("https://books.toscrape.com/", a["href"])) for a in home.select("ul.nav-list ul a")]
site = {}
for name, url in cats:
    while url:
        soup = get(url)
        if soup is None:
            problems.append((url, "category page not reachable"))
            break
        for art in soup.select("article.product_pod"):
            a = art.h3.a
            bid = int(re.search(r"_(\d+)/index\.html$", a["href"]).group(1))
            site[bid] = dict(
                title=a["title"],
                price=float(re.search(r"[\d.]+", art.select_one("p.price_color").get_text()).group()),
                rating=next(RATING[c] for c in art.select_one("p.star-rating")["class"] if c in RATING),
                cat=name,
                link="https://books.toscrape.com/catalogue/" + a["href"].split("../")[-1],
            )
        nxt = soup.select_one("li.next a")
        url = urljoin(url, nxt["href"]) if nxt else None
        time.sleep(0.5)
    print(f"{name}: crawled, {len(site)} books so far", flush=True)

print(f"Site books: {len(site)} | Excel books: {len(xl)}")
for bid in sorted(set(site) ^ set(xl)):
    problems.append((bid, "id only on one side"))
for bid in set(site) & set(xl):
    a, b = site[bid], xl[bid]
    for k in ("title", "cat", "rating", "link"):
        if a[k] != b[k]:
            problems.append((b["link"], f"{k}: site '{a[k]}' vs Excel '{b[k]}'"))
    if abs(a["price"] - b["price"]) > 0.001:
        problems.append((b["link"], f"price: site {a['price']} vs Excel {b['price']}"))

pd.DataFrame(problems, columns=["link", "problem"]).to_csv("recheck_categories_report.csv", index=False)
print(f"PASS 2 DONE: {len(problems)} problems")
sys.exit(0 if not problems else 2)
