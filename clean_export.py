"""kitaplar.csv (raw) -> kitaplar_temiz.csv + kitaplar.xlsx (formatted, dynamic).

Raw file is never modified. Optional bit.ly short links are merged from
qisa_linkler.csv if shorten_links.py was run (otherwise 'Kısa Link' stays empty).
"""
import os
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

RAW, CLEAN, XLSX, SHORT = "kitaplar.csv", "kitaplar_temiz.csv", "kitaplar.xlsx", "qisa_linkler.csv"
COLUMNS = ["Kitap ID", "Kitap Adı", "Fiyat", "Puan", "Stok Durumu", "Stok Adedi", "Kategori", "Link", "Kısa Link"]


def clean():
    df = pd.read_csv(RAW)
    n0 = len(df)
    df = df.drop_duplicates(subset="url").dropna(subset=["kitap_adi", "fiyat", "puan", "kategori", "url"])
    assert len(df) == n0, "unexpected duplicates/missing rows"

    df["kitap_adi"] = df["kitap_adi"].str.replace(r"\s+", " ", regex=True).str.strip()
    df["kategori"] = df["kategori"].str.strip()
    df["stok_durumu"] = df["stok_durumu"].str.strip()
    df["fiyat"] = df["fiyat"].astype(float).round(2)
    df["puan"] = df["puan"].astype(int)
    df["stok_adedi"] = df["stok_adedi"].astype(int)
    df["id"] = df["url"].str.extract(r"_(\d+)/index\.html$")[0].astype(int)
    assert df["id"].is_unique and df["fiyat"].gt(0).all() and df["puan"].between(1, 5).all()

    out = df.rename(columns={
        "id": "Kitap ID", "kitap_adi": "Kitap Adı", "fiyat": "Fiyat", "puan": "Puan",
        "stok_durumu": "Stok Durumu", "stok_adedi": "Stok Adedi", "kategori": "Kategori", "url": "Link",
    })
    out["Kısa Link"] = ""
    if os.path.exists(SHORT):
        short = pd.read_csv(SHORT).set_index("url")["qisa_link"]
        out["Kısa Link"] = out["Link"].map(short).fillna("")
    out = out[COLUMNS].sort_values(["Kategori", "Kitap Adı"]).reset_index(drop=True)
    out.to_csv(CLEAN, index=False, encoding="utf-8-sig")
    return out


HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SUB_FILL = PatternFill("solid", fgColor="DDEBF7")


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 22


def build_xlsx(df):
    wb = Workbook()
    ws = wb.active
    ws.title = "Kitaplar"
    ws.append(COLUMNS)
    for row in df.itertuples(index=False):
        ws.append(list(row))
    last = len(df) + 1

    # real Excel Table: filter buttons, banded rows, grows with new data
    tab = Table(displayName="Kitaplar", ref=f"A1:{get_column_letter(len(COLUMNS))}{last}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)
    style_header(ws, 1, len(COLUMNS))

    for r in range(2, last + 1):
        ws.cell(r, 3).number_format = '0.00'
        for c in (1, 4, 6):
            ws.cell(r, c).alignment = Alignment(horizontal="center")
        link = ws.cell(r, 8)
        link.hyperlink = link.value
        link.font = Font(color="0563C1", underline="single")
        if ws.cell(r, 9).value:
            ws.cell(r, 9).hyperlink = ws.cell(r, 9).value
            ws.cell(r, 9).font = Font(color="0563C1", underline="single")
    for col, w in zip("ABCDEFGHI", (10, 52, 10, 8, 13, 12, 22, 72, 24)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"

    # formulas point at the table, so they update when rows are added/changed
    cs = wb.create_sheet("Kateqoriya özət")
    cs.append(["Kategori", "Kitap Sayı", "Ortalama Fiyat"])
    style_header(cs, 1, 3)
    cats = sorted(df["Kategori"].unique())
    for i, cat in enumerate(cats, start=2):
        cs.cell(i, 1, cat)
        cs.cell(i, 2, f"=COUNTIFS(Kitaplar[Kategori],A{i})")
        cs.cell(i, 3, f"=AVERAGEIFS(Kitaplar[Fiyat],Kitaplar[Kategori],A{i})").number_format = '0.00'
    cs.column_dimensions["A"].width, cs.column_dimensions["B"].width, cs.column_dimensions["C"].width = 24, 12, 16
    cs.freeze_panes = "A2"
    cs.auto_filter.ref = f"A1:C{len(cats) + 1}"

    ps = wb.create_sheet("Puan paylanması")
    ps.append(["Puan", "Kitap Sayı", "Ortalama Fiyat"])
    style_header(ps, 1, 3)
    for p in range(1, 6):
        ps.cell(p + 1, 1, p)
        ps.cell(p + 1, 2, f"=COUNTIFS(Kitaplar[Puan],A{p + 1})")
        ps.cell(p + 1, 3, f"=AVERAGEIFS(Kitaplar[Fiyat],Kitaplar[Puan],A{p + 1})").number_format = '0.00'
    for col, w in zip("ABC", (8, 12, 16)):
        ps.column_dimensions[col].width = w
    wb.save(XLSX)


if __name__ == "__main__":
    data = clean()
    build_xlsx(data)
    print(f"{len(data)} rows -> {CLEAN}, {XLSX}")
